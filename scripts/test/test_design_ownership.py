"""Input adapters, project canvas and template scope share their real production interfaces."""
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(SCRIPTS), str(SCRIPTS / 'orchestrate'), str(SCRIPTS / 'template')]
from canvas_frame import Canvas, project_frame
from validate_svg_layout import validate_file
import workbench_store as store
import ppt_pipeline as pipeline
from project_state import parse_position
from template_library import validate_pack_dir
from PIL import Image
from pptx import Presentation


class DesignOwnershipTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)

    def make_project(self, route='video', width=1080, height=1920):
        source = self.base / 'source.md'
        source.write_text('# A direct brief\nContent supplied without a production package.\n')
        root = self.base / 'project'
        result = subprocess.run([sys.executable, str(SCRIPTS / 'init_svg_project.py'), str(root),
                                 '--source', str(source), '--route', route, '--width', str(width),
                                 '--height', str(height)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        content = root / '_internal/01_content/page_content.json'
        data = json.loads(content.read_text())
        data['pages'] = [{'page_key': 'p1', 'title': 'Direct brief', 'content': 'Direct brief', 'notes': ''}]
        content.write_text(json.dumps(data))
        (root / '_internal/01_content/design_direction.md').write_text('Direct reference. No person window. No audio binding.')
        checkout = store.commit(root, {'op': 'checkout', 'operation_id': 'checkout1', 'page_key': 'p1'})
        candidate = Path(checkout['candidate'])
        candidate.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">'
                             f'<rect id="p1-bg" width="{width}" height="{height}" fill="#ffffff"/>'
                             '<text id="p1-title" x="80" y="180" font-size="48" fill="#111111" font-family="Arial">Direct brief</text></svg>')
        store.commit(root, {'op': 'save', 'operation_id': 'save1', 'page_key': 'p1', 'author': 'model',
                            'candidate': str(candidate), 'base_revision': checkout['base_revision'], 'notes': ''})
        return root

    def test_direct_video_exports_honestly_unbound_without_upstream_skills(self):
        root = self.make_project()
        result = pipeline.export(root)
        self.assertEqual(result['route'], 'video')
        self.assertEqual(result['content_binding'], 'unbound')
        frozen = store.load_snapshot(root, result['snapshot_path'])
        self.assertEqual(frozen['canvas'], {'width': 1080, 'height': 1920})
        self.assertNotIn('visual_contract', frozen['upstream'])

    def test_snapshot_reexport_uses_frozen_canvas_after_live_settings_change(self):
        root = self.make_project()
        first = pipeline.export(root)
        path = root / '_internal/00_project/page_manifest.json'
        data = json.loads(path.read_text()); data['canvas'] = {'width': 1920, 'height': 1080}
        path.write_text(json.dumps(data))
        self.assertEqual(pipeline.export(root, first['snapshot_path'])['snapshot_id'], first['snapshot_id'])
        svg = store.get_page(root, 'p1')['svg_path']
        codes = {i['code'] for i in validate_file(svg, route='video')['issues']}
        self.assertIn('INVALID_CANVAS_SIZE', codes)

    def test_direct_slides_export_preserves_selected_aspect_and_editable_text(self):
        root = self.make_project('slides', 1200, 900)
        result = pipeline.export(root)
        presentation = Presentation(result['output_path'])
        self.assertAlmostEqual(presentation.slide_width / presentation.slide_height, 4/3, places=4)
        texts = [s.text for s in presentation.slides[0].shapes if s.has_text_frame]
        self.assertIn('Direct brief', texts)

    def test_portrait_inspection_accepts_lower_half_and_rejects_overflow(self):
        frame = Canvas(1080, 1920)
        self.assertEqual(parse_position('x=80,y=1700,w=400,h=80', frame)['y'], 1700)
        with self.assertRaises(ValueError):
            parse_position('x=900,y=1700,w=400,h=80', frame)

    def test_portrait_bounds_catch_elements_outside_selected_width(self):
        root = self.make_project()
        path = self.base / 'invalid.svg'
        svg = store.get_page(root, 'p1')['svg']
        path.write_text(svg.replace('x="80"', 'x="1090"'))
        codes = {i['code'] for i in validate_file(path, route='video', frame=project_frame(root))['issues']}
        self.assertIn('TEXT_OUTSIDE_CANVAS', codes)

    @unittest.skipUnless(os.environ.get('PPT_WORKBENCH_BROWSER_TESTS') == '1', 'requires browser renderer')
    def test_portrait_check_renders_full_size_and_inspects_bottom(self):
        root = self.make_project()
        result = pipeline.check(root, ['p1'])['pages'][0]
        self.assertEqual(result['status'], 'pass', result)
        with Image.open(result['png']) as image:
            self.assertEqual(image.size, (1080, 1920))
        path = root / '_internal/00_project/page_manifest.json'
        data = json.loads(path.read_text()); data['canvas'] = {'width': 1920, 'height': 1080}
        path.write_text(json.dumps(data))
        self.assertEqual(pipeline.check(root, ['p1'])['pages'][0]['status'], 'fix')

    def test_legacy_default_and_invalid_project_canvas(self):
        self.assertEqual(project_frame(self.base), Canvas(1920, 1080))
        for values in [(0, 900), (-1, 900), (True, 900), (1200.5, 900)]:
            with self.assertRaises(ValueError): Canvas(*values)

    def test_missing_production_binding_is_not_downgraded_to_standalone(self):
        root = self.make_project()
        head = root / '_internal/06_workbench/head.json'
        data = json.loads(head.read_text()); data['input_kind'] = 'production'
        head.write_text(json.dumps(data))
        with self.assertRaises(store.StoreError) as caught: store.snapshot(root, 'video')
        self.assertEqual(caught.exception.code, 'missing_contract')

    def test_a_single_verified_use_case_is_a_valid_package_but_empty_is_not(self):
        pack = self.base / 'pack'; (pack / 'primitives').mkdir(parents=True); (pack / 'anchors').mkdir()
        (pack / 'SPEC.md').write_text('Cover only; verified for short titles.')
        (pack / 'tokens.css').write_text(':root { --ink: #111111; }')
        svg = '<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" viewBox="0 0 1920 1080"><rect width="1920" height="1080" fill="#fff"/></svg>'
        (pack / 'skyline_shell.svg').write_text(svg)
        primitive = pack / 'primitives/cover.svg'; primitive.write_text(svg)
        Image.new('RGB', (32, 18), 'white').save(pack / 'anchors/reference.png')
        self.assertTrue(validate_pack_dir(pack)['valid'])
        primitive.unlink()
        self.assertFalse(validate_pack_dir(pack)['valid'])


if __name__ == '__main__': unittest.main()
