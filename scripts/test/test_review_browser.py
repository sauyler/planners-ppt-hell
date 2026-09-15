"""Optional browser/interface smoke. All approvals below are synthetic test data."""
import os
import sys
from pathlib import Path
import threading
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parent))
import test_v5 as fixtures
from test_v5 import S
from project_state import *
from ppt_pipeline import export, check
from review_server import ReviewHandler, ThreadingHTTPServer

@unittest.skipUnless(os.environ.get('PPT_BROWSER_TESTS')=='1','set PPT_BROWSER_TESTS=1 for local browser/interface smoke')
class ReviewBrowserTests(unittest.TestCase):
    def test_real_browser_feedback_and_export(self):
        fixture=fixtures.V5Tests();fixture.setUp();root=fixture.root.resolve()
        try:
            fixture.payload()
            class Handler(ReviewHandler):pass
            Handler.project_root=root;Handler.session_id='synthetic-browser-test'
            server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
            thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                browser=p.chromium.launch(headless=True);page=browser.new_page(viewport={'width':1440,'height':900});errors=[]
                page.on('pageerror',lambda error: errors.append(str(error)))
                page.goto(f'http://127.0.0.1:{server.server_port}/review')
                self.assertEqual(page.locator('#rail button').count(),2)
                page.locator('#feedback').fill('Move title');self.assertIn('要求修改',page.locator('#pageStatus').inner_text())
                page.locator('#feedback').fill('');self.assertIn('未处理',page.locator('#pageStatus').inner_text())
                image=root/PNG/'alpha.png'
                page.locator('#newFile').set_input_files(str(image));page.wait_for_selector('.asset')
                page.locator('.asset select').first.select_option('cover')
                self.assertIn('要求修改',page.locator('#pageStatus').inner_text())
                page.get_by_text('移除新增图片',exact=True).click();self.assertEqual(page.locator('.asset').count(),0)
                self.assertIn('未处理',page.locator('#pageStatus').inner_text())
                page.get_by_text('框选问题',exact=True).click();box=page.locator('#preview').bounding_box()
                page.mouse.move(box['x']+box['width']*.1,box['y']+box['height']*.1);page.mouse.down();page.mouse.move(box['x']+box['width']*.4,box['y']+box['height']*.4);page.mouse.up()
                page.locator('#regionNote').fill('Fix region');page.get_by_text('保存标注',exact=True).click();self.assertEqual(page.locator('.annotation').count(),1)
                page.get_by_text('清空框选',exact=True).click()
                page.get_by_text('提交本轮审阅',exact=True).first.click();page.get_by_text('批准未处理页并提交',exact=True).click()
                page.wait_for_function("document.querySelector('#submitMessage').textContent.includes('已保存')")
                self.assertTrue(approved(root));self.assertEqual(errors,[])
                browser.close()
            server.shutdown();server.server_close();thread.join()
            output=export(root)
            self.assertEqual(output['page_count'],2)
            from pptx import Presentation
            prs=Presentation(root/'final_deck.pptx');self.assertTrue(any(s.has_text_frame and s.text for slide in prs.slides for s in slide.shapes))
            self.assertEqual(export(root)['pptx_sha256'],output['pptx_sha256'])
            # Real renderer, cache and stale inspection behavior on an isolated fixture.
            fixture.svg('alpha','New title')
            results=check(root,['alpha']);self.assertEqual(results[0]['status'],'pass')
            results=check(root,['alpha']);self.assertEqual(results[0]['status'],'unchanged')
            import subprocess
            result=subprocess.run([sys.executable,str(S/'render_svg_png.py'),str(root/SVG),str(root/'template-preview-smoke')],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            self.assertTrue(read(root/'template-preview-smoke/png_manifest.json')['all_valid_size'])
        finally:fixture.tearDown()

if __name__=='__main__':unittest.main()
